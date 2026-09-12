import pytest
import pytest_asyncio
import os
from io import BytesIO
from httpx import AsyncClient, ASGITransport
from datetime import datetime, timezone
from motor.motor_asyncio import AsyncIOMotorClient
from motor.core import AgnosticClient

if not hasattr(AgnosticClient, "append_metadata"):
    setattr(AgnosticClient, "append_metadata", lambda self, *args, **kwargs: None)

from beanie import init_beanie
from app.main import app
from app.core.config import settings
from app.models.organization import Organization
from app.models.user import User
from app.models.employee import Employee
from app.models.attrition import AttritionPrediction
from app.models.job import Job
from app.models.candidate import Candidate
from app.models.match import CandidateMatch
from app.core.security import create_access_token, get_password_hash

SAMPLE_RESUME_TEXT = """
John Doe
Senior Software Engineer
Email: john.doe@example.com
Phone: +1-555-123-4567

Summary:
Experienced backend engineer with 3.5 years of experience building Python and FastAPI microservices.

Skills:
Python, FastAPI, SQL, PostgreSQL, MongoDB, Docker, Git, REST API

Education:
B.Tech in Computer Science and Engineering, 2021

Experience:
Software Engineer at TechCorp Inc (2021 - Present)
- Built high-performance FastAPI services and MongoDB databases.
"""

@pytest_asyncio.fixture(autouse=True)
async def setup_test_db():
    client = AsyncIOMotorClient(settings.MONGODB_URL)
    await init_beanie(
        database=client["neurohr_test"],
        document_models=[Organization, User, Employee, AttritionPrediction, Job, Candidate, CandidateMatch]
    )

    await CandidateMatch.find_all().delete()
    await Candidate.find_all().delete()
    await Job.find_all().delete()
    await Employee.find_all().delete()
    await AttritionPrediction.find_all().delete()
    await User.find_all().delete()
    await Organization.find_all().delete()
    yield

@pytest_asyncio.fixture
async def test_org():
    org = Organization(name="Acme Corp", slug="acme-corp")
    await org.insert()
    return org

@pytest_asyncio.fixture
async def test_user(test_org):
    user = User(
        organization_id=test_org,
        email="recruiter@acmecorp.com",
        password_hash=get_password_hash("password123"),
        role="RECRUITER"
    )
    await user.insert()
    return user

@pytest_asyncio.fixture
async def auth_headers(test_user):
    token = create_access_token(str(test_user.id))
    return {"Authorization": f"Bearer {token}"}

@pytest.mark.asyncio
async def test_job_creation_and_retrieval(auth_headers):
    """1. Test creating and retrieving jobs"""
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        payload = {
            "title": "Backend Developer",
            "description": "Build scalable python services",
            "required_skills": ["Python", "FastAPI", "SQL", "MongoDB"],
            "minimum_experience": 2.0,
            "qualification": "B.E / B.Tech / MCA",
            "location": "Coimbatore",
            "employment_type": "Full Time",
            "status": "Open"
        }
        res = await ac.post("/api/v1/recruitment/jobs", json=payload, headers=auth_headers)
        assert res.status_code == 200
        data = res.json()
        assert data["title"] == "Backend Developer"
        assert data["job_id"].startswith("JOB-")
        job_id = data["job_id"]

        # List jobs
        list_res = await ac.get("/api/v1/recruitment/jobs", headers=auth_headers)
        assert list_res.status_code == 200
        jobs = list_res.json()
        assert len(jobs) == 1
        assert jobs[0]["job_id"] == job_id

        # Get single job
        get_res = await ac.get(f"/api/v1/recruitment/jobs/{job_id}", headers=auth_headers)
        assert get_res.status_code == 200
        assert get_res.json()["title"] == "Backend Developer"

@pytest.mark.asyncio
async def test_resume_upload_validation(auth_headers):
    """2. Test invalid file rejection on resume upload"""
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        files = {"file": ("malicious.exe", BytesIO(b"binary content"), "application/octet-stream")}
        res = await ac.post("/api/v1/recruitment/candidates/upload", headers=auth_headers, files=files)
        assert res.status_code == 400
        assert "Invalid file extension" in res.json()["detail"]

@pytest.mark.asyncio
async def test_resume_upload_and_parsing(auth_headers):
    """3. Test valid resume text parsing and candidate creation"""
    # Create sample docx in memory
    import docx
    doc = docx.Document()
    for line in SAMPLE_RESUME_TEXT.strip().split("\n"):
        doc.add_paragraph(line)
    
    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        files = {"file": ("john_doe_resume.docx", bio, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        res = await ac.post("/api/v1/recruitment/candidates/upload", headers=auth_headers, files=files)
        assert res.status_code == 200
        candidate = res.json()
        assert candidate["candidate_id"].startswith("CAND-")
        assert candidate["name"] == "John Doe"
        assert candidate["email"] == "john.doe@example.com"
        assert "Python" in candidate["skills"]
        assert "FastAPI" in candidate["skills"]
        assert candidate["experience"] == 3.5

@pytest.mark.asyncio
async def test_candidate_ranking_calculation(auth_headers):
    """4. Test matching algorithm and candidate ranking calculation"""
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        # Create Job
        job_payload = {
            "title": "Backend Developer",
            "description": "Build python backend services with FastAPI and SQL",
            "required_skills": ["Python", "FastAPI", "SQL", "Docker"],
            "minimum_experience": 2.0,
            "qualification": "B.Tech",
            "location": "Coimbatore",
            "employment_type": "Full Time",
            "status": "Open"
        }
        job_res = await ac.post("/api/v1/recruitment/jobs", json=job_payload, headers=auth_headers)
        job_id = job_res.json()["job_id"]

        # Create docx resume
        import docx
        doc = docx.Document()
        for line in SAMPLE_RESUME_TEXT.strip().split("\n"):
            doc.add_paragraph(line)
        bio = BytesIO()
        doc.save(bio)
        bio.seek(0)

        # Upload Candidate
        files = {"file": ("john_doe_resume.docx", bio, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        await ac.post("/api/v1/recruitment/candidates/upload", headers=auth_headers, files=files)

        # Rank candidates for job
        rank_res = await ac.post(f"/api/v1/recruitment/jobs/{job_id}/rank", headers=auth_headers)
        assert rank_res.status_code == 200
        data = rank_res.json()
        assert data["job_id"] == job_id
        assert data["total_candidates"] == 1
        top_cand = data["rankings"][0]
        assert top_cand["rank"] == 1
        assert top_cand["match_score"] > 50.0
        assert "Python" in top_cand["matched_skills"]
        assert "Docker" in top_cand["matched_skills"] or "Docker" in top_cand["missing_skills"]
        assert "Positives" in top_cand["explanation"]

@pytest.mark.asyncio
async def test_organization_isolation(auth_headers, test_org):
    """5. Verify multi-tenant organization isolation for recruitment data"""
    org2 = Organization(name="Competitor Inc", slug="competitor-inc")
    await org2.insert()
    user2 = User(organization_id=org2, email="recruiter2@competitor.com", password_hash=get_password_hash("pass"), role="RECRUITER")
    await user2.insert()
    headers2 = {"Authorization": f"Bearer {create_access_token(str(user2.id))}"}

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        # Org 1 creates job
        job_payload = {"title": "Acme Special Role", "description": "Confidential", "required_skills": ["Python"]}
        job_res = await ac.post("/api/v1/recruitment/jobs", json=job_payload, headers=auth_headers)
        job_id = job_res.json()["job_id"]

        # Org 2 attempts to list jobs -> should see 0 jobs
        list2 = await ac.get("/api/v1/recruitment/jobs", headers=headers2)
        assert list2.status_code == 200
        assert len(list2.json()) == 0

        # Org 2 attempts to access Org 1's job -> should receive 404
        get2 = await ac.get(f"/api/v1/recruitment/jobs/{job_id}", headers=headers2)
        assert get2.status_code == 404
